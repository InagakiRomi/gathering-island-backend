import random
import sys
from pathlib import Path

import pandas as pd

_DATA_DIR = Path(__file__).resolve().parent / ".." / "data"
GATHERING_PATH = _DATA_DIR / "Gathering.xlsx"
PARTICIPANT_OUTPUT_PATH = _DATA_DIR / "Participant_filled.xlsx"

# 參與者 userId 僅限此範圍（2～39）
VALID_USER_ID_MIN = 2
VALID_USER_ID_MAX = 39
VALID_USER_IDS = list(range(VALID_USER_ID_MIN, VALID_USER_ID_MAX + 1))

# 與 Participant.xlsx 現有格式一致：第一行中文、第二行英文
PARTICIPANT_ROW0 = ["對應gatheringId", "對應userId", "加入時間"]
PARTICIPANT_HEADER = ["gathering", "user", "joinedAt"]


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    return df.rename(columns=lambda c: c.strip() if isinstance(c, str) else c)


def main() -> None:
    if not GATHERING_PATH.exists():
        print(f"錯誤：找不到 {GATHERING_PATH}")
        input("按 Enter 結束...")
        return

    df_g = _normalize_columns(pd.read_excel(GATHERING_PATH, header=1))
    for col in ["createdAt", "deadline"]:
        if col in df_g.columns:
            df_g[col] = pd.to_datetime(df_g[col], errors="coerce")

    # 參與者僅能從 userId 2～39 中選，且同一活動內不重複（random.sample 已保證不重複）
    other_users_pool = VALID_USER_IDS
    if len(other_users_pool) < 2:
        print("參與者 userId 範圍 2～39 至少需 2 人，無法產生參與者。")
        input("按 Enter 結束...")
        return

    rows = []
    random.seed(42)
    for _, g in df_g.iterrows():
        gid = int(g["id"])
        creator = int(g["userId"])
        cap = int(g["participantNumbers"]) if pd.notna(g["participantNumbers"]) else 0
        created = g.get("createdAt")
        deadline = g.get("deadline")
        if pd.isna(created) or pd.isna(deadline) or created >= deadline:
            continue
        # 僅能選 2～39 且排除該活動建立者；同一活動內每人只出現一次（sample 不重複）
        other_users = [u for u in other_users_pool if u != creator]
        max_n = min(max(0, cap - 1), len(other_users))
        if max_n <= 0:
            continue
        # 參加人數隨機多寡：0 ～ max_n，讓每場活動人數差異大
        n = random.randint(0, max_n)
        if n == 0:
            continue
        chosen = random.sample(other_users, n)  # 同一活動內參加者不重複
        start_ts = created.timestamp()
        end_ts = deadline.timestamp()
        for i, uid in enumerate(chosen):
            # 在 [createdAt, deadline] 內均勻分配
            t = start_ts + (end_ts - start_ts) * (i + 1) / (n + 1)
            joined = pd.Timestamp.utcfromtimestamp(t).tz_localize(None)
            if hasattr(joined, "strftime"):
                joined_str = joined.strftime("%Y-%m-%d %H:%M:%S")
            else:
                joined_str = str(joined)
            rows.append({"gathering": gid, "user": uid, "joinedAt": joined_str})

    df_p = pd.DataFrame(rows)
    if df_p.empty:
        print("未產生任何參與者資料。")
        input("按 Enter 結束...")
        return

    # 依活動、加入時間排序
    df_p = df_p.sort_values(by=["gathering", "joinedAt"]).reset_index(drop=True)

    # 檢查：每場活動報名人數（參與者數）不得超過可報名額 = participantNumbers - 1（建立者佔 1）
    cap_by_gid = df_g.set_index("id")["participantNumbers"].astype(int).to_dict()
    over = []
    for gid, count in df_p.groupby("gathering").size().items():
        cap = cap_by_gid.get(gid)
        if cap is not None and count > max(0, cap - 1):
            over.append((gid, count, cap))
    if over:
        print("錯誤：以下活動報名人數超過可報名額（participantNumbers - 1）：", over)
        input("按 Enter 結束...")
        sys.exit(1)
    print("檢查通過：所有活動報名人數皆未超過上限。")

    # 寫入「新檔案」，不覆蓋原 Participant.xlsx
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for col, val in enumerate(PARTICIPANT_ROW0, 1):
        ws.cell(row=1, column=col, value=val)
    for col, val in enumerate(PARTICIPANT_HEADER, 1):
        ws.cell(row=2, column=col, value=val)
    for r, row in df_p.iterrows():
        for c, val in enumerate(row, 1):
            ws.cell(row=int(r) + 3, column=c, value=val)
    wb.save(PARTICIPANT_OUTPUT_PATH)

    print(f"已寫入 {len(df_p)} 筆參與者至 {PARTICIPANT_OUTPUT_PATH}（未覆蓋原 Participant.xlsx）")
    print(f"活動數：{df_p['gathering'].nunique()}，參與者（user）數：{df_p['user'].nunique()}（userId 僅 2～39）")
    # 每場人數分佈
    counts = df_p.groupby("gathering").size()
    print(f"每場報名人數：最少 {int(counts.min())}，最多 {int(counts.max())}，平均 {counts.mean():1.1f}")
    input("按 Enter 結束...")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        input("按 Enter 結束...")
        raise
